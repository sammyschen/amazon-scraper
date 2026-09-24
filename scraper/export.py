"""
Export layer: rows -> the approved workbook (Products + Info sheets), used by every mode.
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import LIST_PRICE_COL, PRICE_COL, PRICE_FORMAT

# The approved Products schema - do not rename, reorder or insert columns.
# The two price columns are named after the marketplace currency (price_usd, list_price_usd).
COLUMNS = [
    "rank", "asin", "brand", "title", "subtitle", "condition",
    PRICE_COL, LIST_PRICE_COL, "list_price_type", "discount_pct",
    "rating", "ratings_count", "bought_past_month", "stock", "badge", "deal",
    "advertised", "delivery", "other_offers", "found_via", "page",
    "url", "image_url", "scraped_at",
]
WIDTHS = {"title": 60, "subtitle": 45, "url": 38, "image_url": 45, "stock": 22,
          "delivery": 28, "other_offers": 40, "badge": 32, "found_via": 20}


def export_results(results, metadata, path, csv=False):
    """Write rows (dicts with the COLUMNS keys) and run metadata ({field: value}) to
    path.xlsx (+ path.csv). Returns the list of files written."""
    path = Path(path)
    df = pd.DataFrame(results).reindex(columns=COLUMNS)
    df["advertised"] = df["advertised"].map({True: "Yes", False: "No"})

    xlsx = path.with_suffix(".xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Products", index=False)
        pd.DataFrame(list(metadata.items()), columns=["field", "value"]).to_excel(
            xw, sheet_name="Info", index=False)

        ws = xw.sheets["Products"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="232F3E")
            cell.alignment = Alignment(vertical="center")
        for idx, name in enumerate(df.columns, start=1):
            col = get_column_letter(idx)
            ws.column_dimensions[col].width = WIDTHS.get(name, max(12, len(name) + 2))
            for cell in ws[col][1:]:
                if name in (PRICE_COL, LIST_PRICE_COL):
                    cell.number_format = PRICE_FORMAT
                elif name == "ratings_count":
                    cell.number_format = "#,##0"
                elif name == "url" and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
        xw.sheets["Info"].column_dimensions["A"].width = 14
        xw.sheets["Info"].column_dimensions["B"].width = 100

    written = [xlsx]
    if csv:
        df.to_csv(path.with_suffix(".csv"), index=False, encoding="utf-8-sig")
        written.append(path.with_suffix(".csv"))
    return written
